import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import chip_report  # noqa: E402
import events  # noqa: E402
import extras  # noqa: E402
import finmind  # noqa: E402
import notify  # noqa: E402
import pages  # noqa: E402
import weekly  # noqa: E402
from analysis import apply_streaks, build_report, build_stocks  # noqa: E402
from notify import build_summary  # noqa: E402
from outputs import render_html, render_markdown, write_excel  # noqa: E402

STOCK_INFO = [
    {"stock_id": "2330", "stock_name": "台積電", "type": "twse", "industry_category": "半導體業"},
    {"stock_id": "2330", "stock_name": "台積電", "type": "twse", "industry_category": "電子工業"},
    {"stock_id": "2317", "stock_name": "鴻海", "type": "twse", "industry_category": "其他電子業"},
    {"stock_id": "6488", "stock_name": "環球晶", "type": "tpex", "industry_category": "半導體業"},
    {"stock_id": "0050", "stock_name": "元大台灣50", "type": "twse", "industry_category": "ETF"},
    {"stock_id": "00878", "stock_name": "國泰永續高股息", "type": "twse", "industry_category": "ETF"},
    {"stock_id": "0056", "stock_name": "元大高股息", "type": "twse", "industry_category": "ETF"},
]

INSTITUTIONAL = [
    {"stock_id": "2330", "name": "Foreign_Investor", "buy": 5_000_000, "sell": 2_000_000},
    {"stock_id": "2330", "name": "Foreign_Dealer_Self", "buy": 1_000, "sell": 0},
    {"stock_id": "2330", "name": "Investment_Trust", "buy": 300_000, "sell": 100_000},
    {"stock_id": "2330", "name": "Dealer_self", "buy": 50_000, "sell": 10_000},
    {"stock_id": "2317", "name": "Foreign_Investor", "buy": 100_000, "sell": 300_000},
    {"stock_id": "2317", "name": "Investment_Trust", "buy": 0, "sell": 50_000},
    {"stock_id": "6488", "name": "Foreign_Investor", "buy": 100_000, "sell": 900_000},
    {"stock_id": "0050", "name": "Foreign_Investor", "buy": 9_999_999, "sell": 0},
    {"stock_id": "00878", "name": "Foreign_Investor", "buy": 2_000_000, "sell": 0},
    {"stock_id": "00878", "name": "Dealer_Hedging", "buy": 0, "sell": 500_000},
    {"stock_id": "0056", "name": "Investment_Trust", "buy": 0, "sell": 100_000},
]

PRICES = [
    {"stock_id": "2330", "close": 1000.0, "spread": 20.0, "Trading_Volume": 30_000_000},
    {"stock_id": "2317", "close": 200.0, "spread": -4.0, "Trading_Volume": 20_000_000},
    # 6488 無股價，測試缺值處理
]


_patchers = []


def setUpModule():
    # 測試不連外網：行事曆的 HTTP 請求一律失敗（FOMC 會改用內建日程），且不帶 API 金鑰
    def offline(*args, **kwargs):
        raise events.requests.ConnectionError("offline")

    _patchers.append(mock.patch.object(events.requests, "get", side_effect=offline))
    _patchers.append(mock.patch.dict(os.environ, {"FRED_API_KEY": "", "ALPHAVANTAGE_API_KEY": ""}))
    for p in _patchers:
        p.start()


def tearDownModule():
    for p in reversed(_patchers):
        p.stop()


def _inst(stock_id, name, net):
    return {"stock_id": stock_id, "name": name,
            "buy": max(net, 0), "sell": max(-net, 0)}


# 2026-09-25（五）之前的歷史資料；09-22（二）當作休市日沒有資料
HISTORY = {
    "2026-09-24": [_inst("2330", "Foreign_Investor", 1_000_000),
                   _inst("2330", "Investment_Trust", 50_000),
                   _inst("2317", "Foreign_Investor", -100_000)],
    "2026-09-23": [_inst("2330", "Foreign_Investor", 2_000_000),
                   _inst("2330", "Investment_Trust", -10_000),
                   _inst("2317", "Foreign_Investor", 300_000)],
    "2026-09-21": [_inst("2330", "Foreign_Investor", 500_000),
                   _inst("2330", "Foreign_Dealer_Self", -1_000)],
    "2026-09-18": [_inst("2330", "Foreign_Investor", -700_000)],
}


def _oi(day, investor, long_oi, short_oi, **extra):
    return {"date": day, "institutional_investors": investor,
            "long_open_interest_balance_volume": long_oi,
            "short_open_interest_balance_volume": short_oi, **extra}


FUTURES = [
    _oi("2026-09-24", "外資", 21_000, 49_000),
    _oi("2026-09-24", "投信", 29_000, 5_000),
    _oi("2026-09-25", "外資", 20_000, 50_000),
    _oi("2026-09-25", "投信", 30_000, 5_000),
    _oi("2026-09-25", "自營商", 10_000, 12_000),
]

OPTION_INST = [
    _oi("2026-09-24", "外資", 28_000, 20_000, call_put="買權"),
    _oi("2026-09-25", "外資", 30_000, 20_000, call_put="買權"),
    _oi("2026-09-25", "外資", 15_000, 25_000, call_put="賣權"),
    _oi("2026-09-25", "投信", 1, 1, call_put="買權"),
]

OPTION_DAILY = {
    "2026-09-25": [
        {"call_put": "call", "open_interest": 60_000, "trading_session": "position"},
        {"call_put": "call", "open_interest": 40_000, "trading_session": "position"},
        {"call_put": "put", "open_interest": 120_000, "trading_session": "position"},
        {"call_put": "put", "open_interest": 999, "trading_session": "after_market"},
    ],
    "2026-09-24": [
        {"call_put": "call", "open_interest": 100_000, "trading_session": "position"},
        {"call_put": "put", "open_interest": 100_000, "trading_session": "position"},
    ],
}

MARGIN = [
    {"stock_id": "2330", "MarginPurchaseTodayBalance": 20_000, "MarginPurchaseYesterdayBalance": 19_000,
     "ShortSaleTodayBalance": 400, "ShortSaleYesterdayBalance": 500},
    {"stock_id": "2317", "MarginPurchaseTodayBalance": 30_000, "MarginPurchaseYesterdayBalance": 31_000,
     "ShortSaleTodayBalance": 1_500, "ShortSaleYesterdayBalance": 1_000},
]


def _holding(stock_id, big_pct, people):
    return [
        {"stock_id": stock_id, "HoldingSharesLevel": "800,001-1,000,000", "percent": 5.0, "people": 10},
        {"stock_id": stock_id, "HoldingSharesLevel": "more than 1,000,001", "percent": big_pct, "people": 100},
        {"stock_id": stock_id, "HoldingSharesLevel": "total", "percent": 100.0, "people": people},
    ]


HOLDING = {
    "2026-09-25": _holding("2330", 80.5, 1_000_000) + _holding("2317", 40.0, 500_000),
    "2026-09-18": _holding("2330", 80.0, 1_010_000) + _holding("2317", 41.0, 490_000),
}


def fake_fetch(token, dataset, **params):
    if dataset == "TaiwanStockInfo":
        return STOCK_INFO
    day = params.get("start_date")
    if dataset == "TaiwanFuturesInstitutionalInvestors":
        return FUTURES
    if dataset == "TaiwanOptionInstitutionalInvestors":
        return OPTION_INST
    if dataset == "TaiwanOptionDaily":
        return OPTION_DAILY.get(day, [])
    if dataset == "TaiwanStockHoldingSharesPer":
        return HOLDING.get(day, [])
    if dataset == "TaiwanStockMarginPurchaseShortSale":
        return MARGIN if day == "2026-09-25" else []
    if dataset == "TaiwanStockInstitutionalInvestorsBuySell" and day in HISTORY:
        return HISTORY[day]
    if day != "2026-09-25":
        return []
    if dataset == "TaiwanStockInstitutionalInvestorsBuySell":
        return INSTITUTIONAL
    if dataset == "TaiwanStockPrice":
        return PRICES
    raise AssertionError(dataset)


class ReportTest(unittest.TestCase):
    def setUp(self):
        patcher = mock.patch.object(finmind, "fetch_dataset", fake_fetch)
        patcher.start()
        self.addCleanup(patcher.stop)
        day = date(2026, 9, 25)
        stocks = build_stocks(
            finmind.load_institutional("t", day),
            finmind.load_prices("t", day),
            finmind.load_stock_info("t"),
        )
        self.stocks = stocks
        self.report = build_report(day, stocks, top=20)
        self.by_key = {s.key: s for s in self.report.sections}

    def test_specific_industry_preferred(self):
        rows = [
            {"stock_id": "2308", "stock_name": "台達電", "type": "twse", "industry_category": "電子工業"},
            {"stock_id": "2308", "stock_name": "台達電", "type": "twse", "industry_category": "電子零組件業"},
            {"stock_id": "2317", "stock_name": "鴻海", "type": "twse", "industry_category": "其他電子業"},
            {"stock_id": "2317", "stock_name": "鴻海", "type": "twse", "industry_category": "電子工業"},
        ]
        with mock.patch.object(finmind, "fetch_dataset", return_value=rows):
            info = finmind.load_stock_info("t")
        self.assertEqual(info["2308"]["industry"], "電子零組件業")
        self.assertEqual(info["2317"]["industry"], "其他電子業")

    def test_etf_excluded_and_markets(self):
        self.assertNotIn("0050", self.stocks)
        self.assertEqual(self.stocks["6488"].market, "上櫃")
        self.assertEqual(self.stocks["2330"].industry, "半導體業")

    def test_price_fields(self):
        tsmc = self.stocks["2330"]
        self.assertAlmostEqual(tsmc.change_pct, 20 / 980 * 100)
        # 外資 +3,001 張 × 1000 元 = 30.01 億
        self.assertAlmostEqual(tsmc.amount_yi(tsmc.net("foreign")), 30.01)
        self.assertAlmostEqual(tsmc.volume_pct(tsmc.net("foreign")), 3_001_000 / 30_000_000 * 100)
        self.assertIsNone(self.stocks["6488"].amount_yi(1000))

    def test_rankings(self):
        ids = lambda key: [r[1] for r in self.by_key[key].rows]  # noqa: E731
        self.assertEqual(ids("foreign_buy"), ["2330"])
        self.assertEqual(ids("foreign_sell"), ["6488", "2317"])
        self.assertEqual(ids("trust_buy"), ["2330"])
        self.assertEqual(ids("trust_sell"), ["2317"])
        self.assertEqual(ids("sync_buy"), ["2330"])

    def test_rank_by_amount(self):
        report = build_report(date(2026, 9, 25), self.stocks, top=20, rank_by="amount")
        sell = next(s for s in report.sections if s.key == "foreign_sell")
        # 6488 無股價，金額視為 0，排在鴻海之後
        self.assertEqual([r[1] for r in sell.rows], ["2317", "6488"])

    def test_industry_and_overview(self):
        industry = {r[0]: r for r in self.by_key["industry"].rows}
        self.assertAlmostEqual(industry["半導體業"][1], 30.01)
        self.assertAlmostEqual(industry["半導體業"][2], 2.0)
        self.assertEqual(self.by_key["industry"].rows[0][0], "半導體業")
        overview = {r[0]: r for r in self.report.overview.rows}
        self.assertEqual(overview["外資"][1], 3001 - 200 - 800)
        self.assertEqual(overview["自營商"][1], 40)

    def test_outputs(self):
        md = render_markdown(self.report)
        self.assertIn("## 外資投信同步買超", md)
        self.assertIn("+30.01", md)
        page = render_html(self.report)
        self.assertIn('class="num up"', page)
        self.assertIn('class="num down"', page)
        summary = build_summary(self.report)
        self.assertIn("1. 2330 台積電 +3,001 張（+30.01 億）", summary)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.xlsx"
            write_excel(self.report, path)
            from openpyxl import load_workbook

            wb = load_workbook(path)
            self.assertIn("全部個股", wb.sheetnames)
            self.assertEqual(wb["外資買超"]["B2"].value, "2330")


class StreakTest(unittest.TestCase):
    def setUp(self):
        patcher = mock.patch.object(finmind, "fetch_dataset", fake_fetch)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.day = date(2026, 9, 25)
        self.stocks = build_stocks(
            finmind.load_institutional("t", self.day),
            finmind.load_prices("t", self.day),
            finmind.load_stock_info("t"),
        )
        self.history = finmind.load_history("t", self.day, 9)
        apply_streaks(self.stocks, self.history)

    def test_history_skips_weekends_and_holidays(self):
        self.assertEqual(
            [d.isoformat() for d, _ in self.history],
            ["2026-09-24", "2026-09-23", "2026-09-21", "2026-09-18"],
        )

    def test_streak_counts(self):
        tsmc, hon_hai, gw = self.stocks["2330"], self.stocks["2317"], self.stocks["6488"]
        # 外資：今日＋09-24＋09-23＋09-21（含外資自營商 -1,000 股仍為買超），09-18 賣超中斷
        self.assertEqual(tsmc.streak["foreign"], 4)
        self.assertEqual(tsmc.streak_net["foreign"], 3_001_000 + 1_000_000 + 2_000_000 + 499_000)
        self.assertEqual(tsmc.streak["trust"], 2)
        self.assertEqual(hon_hai.streak["foreign"], -2)
        self.assertEqual(gw.streak["foreign"], -1)  # 09-24 無資料，中斷
        self.assertEqual(gw.streak["trust"], 0)  # 今日無投信買賣

    def test_streak_sections(self):
        report = build_report(self.day, self.stocks, top=20,
                              streak_window=len(self.history) + 1, streak_min=3)
        by_key = {s.key: s for s in report.sections}
        foreign = by_key["foreign_streak"]
        self.assertEqual([r[1] for r in foreign.rows], ["2330"])
        row = foreign.rows[0]
        self.assertEqual(row[foreign.col("連買天數")], 4)
        self.assertEqual(row[foreign.col("累計買超(張)")], 6500)
        self.assertEqual(by_key["trust_streak"].rows, [])  # 投信只連買 2 天
        sell = by_key["foreign_sell"]
        self.assertEqual(
            {r[1]: r[sell.col("連賣天數")] for r in sell.rows}, {"2317": 2, "6488": 1}
        )
        summary = build_summary(report, report_url="https://example.com/2026-09-25.html")
        self.assertIn("1. 2330 台積電 連 4 天（累計 +6,500 張）", summary)
        self.assertTrue(summary.endswith("完整報告：https://example.com/2026-09-25.html"))

    def test_no_streak_keeps_original_sections(self):
        report = build_report(self.day, self.stocks, top=20)
        keys = [s.key for s in report.sections]
        self.assertNotIn("foreign_streak", keys)
        self.assertNotIn("連買天數", [c.header for c in report.sections[0].columns])


class ExtrasTest(unittest.TestCase):
    def setUp(self):
        patcher = mock.patch.object(finmind, "fetch_dataset", fake_fetch)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.day = date(2026, 9, 25)
        self.stocks = build_stocks(
            finmind.load_institutional("t", self.day),
            finmind.load_prices("t", self.day),
            finmind.load_stock_info("t", ["00878", "0056"]),
        )

    def test_etf_section(self):
        etfs = {k: v for k, v in self.stocks.items() if k in ("00878", "0056")}
        section = extras.build_etf_section(etfs, ["0056", "00878"])
        self.assertEqual([r[0] for r in section.rows], ["00878", "0056"])
        self.assertEqual(section.rows[0][section.col("三大法人(張)")], 1500)
        self.assertEqual(section.rows[0][section.col("自營商(張)")], -500)
        self.assertEqual(self.stocks["00878"].industry, "ETF")
        self.assertNotIn("0050", self.stocks)

    def test_futures_section(self):
        section = chip_report.futures_section("t", self.day)
        rows = {r[0]: r for r in section.rows}
        self.assertEqual(rows["外資台指期"][1:], [-30_000, -2_000, 20_000, 50_000])
        self.assertEqual(rows["投信台指期"][1:3], [25_000, 1_000])
        self.assertIsNone(rows["自營商台指期"][2])  # 前一日無資料
        self.assertEqual(rows["外資台指買權"][1:3], [10_000, 2_000])
        self.assertEqual(rows["外資台指賣權"][1:3], [-10_000, None])
        self.assertEqual(section.meta["pc_ratio"], (120.0, 100.0))
        self.assertIn("Put/Call Ratio：120.00%（前一交易日 100.00%）", section.note)

    def test_futures_without_call_put_column(self):
        rows = [dict(r) for r in OPTION_INST]
        for r in rows:
            r.pop("call_put")
        section = extras.build_futures_section(self.day, FUTURES, rows)
        self.assertEqual([r[0] for r in section.rows], ["外資台指期", "投信台指期", "自營商台指期"])

    def test_margin_sections(self):
        up, down, short = extras.build_margin_sections(
            list(self.stocks.values()), finmind.load_margin("t", self.day), 20
        )
        self.assertEqual([(r[1], r[3]) for r in up.rows], [("2330", 1000)])
        self.assertEqual([(r[1], r[3]) for r in down.rows], [("2317", -1000)])
        self.assertEqual([(r[1], r[5]) for r in short.rows], [("2317", 500)])
        self.assertAlmostEqual(short.rows[0][short.col("券資比")], 5.0)
        self.assertIn("融資 +0 張、融券 +400 張", up.note)

    def test_holder_sections_skip_abnormal_rows(self):
        weeks = [
            (date(2026, 9, 25), HOLDING["2026-09-25"] + _holding("2317", 100.0, 1)),
            (date(2026, 9, 18), HOLDING["2026-09-18"]),
        ]
        # 2317 本週出現股東人數 1 人的異常資料，應被排除
        weeks[0] = (weeks[0][0], [r for r in weeks[0][1]
                                  if not (r["stock_id"] == "2317" and r["people"] == 500_000)])
        up, down = extras.build_holder_sections(list(self.stocks.values()), weeks, 20)
        self.assertEqual([r[1] for r in up.rows + down.rows], ["2330"])

    def test_holder_sections(self):
        weeks = finmind.find_holding_weeks("t", date(2026, 9, 29))
        self.assertEqual([d.isoformat() for d, _ in weeks], ["2026-09-25", "2026-09-18"])
        up, down = extras.build_holder_sections(list(self.stocks.values()), weeks, 20)
        self.assertEqual(up.rows[0][1], "2330")
        self.assertAlmostEqual(up.rows[0][up.col("週增減(百分點)")], 0.5)
        self.assertEqual(up.rows[0][up.col("人數增減")], -10_000)
        self.assertEqual(down.rows[0][1], "2317")
        self.assertIn("集保資料日期 2026-09-25", up.note)


class WeeklyTest(unittest.TestCase):
    def setUp(self):
        patcher = mock.patch.object(finmind, "fetch_dataset", fake_fetch)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.day = date(2026, 9, 25)
        inst = finmind.load_institutional("t", self.day)
        self.stocks = build_stocks(inst, finmind.load_prices("t", self.day), finmind.load_stock_info("t"))
        self.history = finmind.load_history("t", self.day, 9)
        self.days = weekly.week_days(self.day, inst, self.history)
        prev = {"2330": {"close": 950.0}}
        self.report = weekly.build_weekly_report(self.day, self.days, self.stocks, 20, prev)
        self.by_key = {s.key: s for s in self.report.sections}

    def test_week_days_and_previous_day(self):
        self.assertEqual([d.isoformat() for d, _ in self.days],
                         ["2026-09-21", "2026-09-23", "2026-09-24", "2026-09-25"])
        self.assertEqual(weekly.previous_trading_day(self.day, self.history), date(2026, 9, 18))

    def test_weekly_rankings(self):
        buy = self.by_key["weekly_foreign_buy"]
        row = buy.rows[0]
        # 台積電本週外資：09-21 +499,000、09-23 +2,000,000、09-24 +1,000,000、09-25 +3,001,000 股
        self.assertEqual(row[1], "2330")
        self.assertEqual(row[buy.col("週買賣超(張)")], 6500)
        self.assertEqual(row[buy.col("買超天數")], 4)
        self.assertAlmostEqual(row[buy.col("週漲跌幅")], 50 / 950 * 100)
        sell = self.by_key["weekly_foreign_sell"]
        self.assertEqual([r[1] for r in sell.rows], ["6488"])  # 鴻海本週 +300 -100 -200 = 0
        self.assertEqual(sell.rows[0][sell.col("賣超天數")], 1)
        self.assertIsNone(sell.rows[0][sell.col("週漲跌幅")])  # 無上週股價

    def test_weekly_overview_and_daily(self):
        overview = {r[0]: r for r in self.report.overview.rows}
        self.assertEqual(overview["外資"][1], 6500 + 0 - 800)
        daily = self.by_key["weekly_daily"]
        self.assertEqual([r[0] for r in daily.rows], ["09/21（一）", "09/23（三）", "09/24（四）", "09/25（五）"])
        self.assertEqual(self.report.title, "台股每週籌碼週報 09/21～09/25")

    def test_weekly_summary(self):
        text = weekly.build_weekly_summary(self.report, report_url="https://x/weekly-2026-09-25.html")
        self.assertIn("1. 2330 台積電 +6,500 張（買超 4 天）", text)
        self.assertIn("1. 6488 環球晶 -800 張（賣超 1 天）", text)
        self.assertIn("【資金流入產業】\n半導體業", text)
        self.assertTrue(text.endswith("完整週報：https://x/weekly-2026-09-25.html"))

    def test_not_generated_on_other_days_unless_forced(self):
        def run(argv, today):
            with tempfile.TemporaryDirectory() as tmp, \
                    mock.patch.object(chip_report, "taipei_today", return_value=today), \
                    mock.patch.dict(os.environ, {"FINMIND_TOKEN": "t"}), \
                    mock.patch.object(sys, "argv", ["chip_report.py", "--output-dir", tmp,
                                                    "--formats", "md"] + argv), \
                    mock.patch("sys.stdout"):
                self.assertEqual(chip_report.main(), 0)
                return sorted(p.name for p in Path(tmp).iterdir())

        self.assertEqual(run(["--date", "2026-09-24"], date(2026, 9, 24)), ["chip_report_2026-09-24.md"])
        self.assertIn("weekly_report_2026-09-24.md",
                      run(["--date", "2026-09-24", "--weekly"], date(2026, 9, 24)))
        self.assertEqual(run(["--no-weekly"], date(2026, 9, 29)), ["chip_report_2026-09-25.md"])


TW = events.TAIPEI


class EventsTest(unittest.TestCase):
    def test_rule_dates(self):
        self.assertEqual(events.nth_weekday(2026, 10, 4, 3), date(2026, 10, 16))
        self.assertEqual(events.nth_weekday(2026, 5, 0, -1), date(2026, 5, 25))
        self.assertEqual(events.easter(2026), date(2026, 4, 5))
        h = events.us_market_holidays(2026)
        self.assertEqual(h[date(2026, 4, 3)], "耶穌受難日")
        self.assertEqual(h[date(2026, 11, 26)], "感恩節")
        self.assertEqual(h[date(2026, 7, 3)], "獨立紀念日")  # 7/4 週六，提前到週五
        self.assertNotIn(date(2021, 12, 31), events.us_market_holidays(2022))  # 元旦週六不補假
        closes = events.us_early_closes(2026)
        self.assertEqual(set(closes), {date(2026, 11, 27), date(2026, 12, 24)})  # 7/3 已休市

    def test_rule_events(self):
        found = {(e.day, e.title) for e in events.rule_events(date(2026, 10, 1), date(2026, 12, 31))}
        self.assertIn((date(2026, 10, 16), "美股選擇權月到期"), found)
        self.assertIn((date(2026, 12, 18), "美股四巫日（季度選擇權、期貨到期）"), found)
        self.assertIn((date(2026, 10, 21), "台指期／台指選擇權月結算"), found)
        # 2025 年四月第三個週五是耶穌受難日，選擇權改在週四到期
        april = [e.day for e in events.rule_events(date(2025, 4, 1), date(2025, 4, 30))
                 if e.title == "美股選擇權月到期"]
        self.assertEqual(april, [date(2025, 4, 17)])

    def test_timezone_conversion(self):
        cpi_summer = events._timed_event(date(2026, 10, 14), events.FRED_RELEASE_TIME, "CPI", 3)
        cpi_winter = events._timed_event(date(2026, 12, 10), events.FRED_RELEASE_TIME, "CPI", 3)
        self.assertEqual((cpi_summer.day, cpi_summer.time_text), (date(2026, 10, 14), "20:30"))
        self.assertEqual((cpi_winter.day, cpi_winter.time_text), (date(2026, 12, 10), "21:30"))
        fomc = events.fomc_events(date(2026, 10, 1), date(2026, 12, 31), fetch=False)
        self.assertEqual([(e.day, e.time_text) for e in fomc],
                         [(date(2026, 10, 29), "02:00"), (date(2026, 12, 10), "03:00")])
        self.assertIn("點陣圖", fomc[1].title)
        self.assertEqual(fomc[0].note, "內建日程，以 Fed 官網為準")

    def test_parse_fomc_calendar(self):
        html = """<h4>2027 FOMC Meetings</h4>
        <div class="fomc-meeting__month col-xs-5"><strong>January</strong></div>
        <div class="fomc-meeting__date col-xs-4">26-27</div>
        <div class="fomc-meeting__month col-xs-5"><strong>Apr/May</strong></div>
        <div class="fomc-meeting__date col-xs-4">30-1*</div>
        <div class="fomc-meeting__month col-xs-5"><strong>August</strong></div>
        <div class="fomc-meeting__date col-xs-4">16 (notation vote)</div>
        <h4>2026 FOMC Meetings</h4>
        <div class="fomc-meeting__month col-xs-5"><strong>December</strong></div>
        <div class="fomc-meeting__date col-xs-4">8-9*</div>"""
        self.assertEqual(events.parse_fomc_calendar(html), [
            (date(2027, 1, 27), False), (date(2027, 5, 1), True), (date(2026, 12, 9), True),
        ])

    def test_parse_earnings_csv(self):
        text = (
            "symbol,name,reportDate,fiscalDateEnding,estimate,currency,timeOfTheDay\n"
            "NVDA,NVIDIA,2026-11-18,2026-10-31,1.2,USD,post-market\n"
            "JPM,JPMorgan,2026-10-13,2026-09-30,4.1,USD,pre-market\n"
            "XYZ,Other,2026-10-13,2026-09-30,1,USD,post-market\n"
        )
        result = {e.title: e for e in events.parse_earnings_csv(text)}
        self.assertEqual(set(result), {"NVDA 輝達 財報", "JPM 摩根大通 財報"})
        nvda = result["NVDA 輝達 財報"]
        self.assertEqual((nvda.day, nvda.time_text, nvda.stars), (date(2026, 11, 19), "清晨", 3))
        self.assertEqual(result["JPM 摩根大通 財報"].time_text, "晚上")
        no_timing = events.parse_earnings_csv(
            "symbol,name,reportDate,fiscalDateEnding,estimate,currency\nAMD,AMD,2026-10-27,,,USD\n")
        self.assertEqual((no_timing[0].day, no_timing[0].time_text), (date(2026, 10, 27), ""))

    def test_fred_events(self):
        payload = {"release_dates": [
            {"release_id": 10, "release_name": "Consumer Price Index", "date": "2026-10-14"},
            {"release_id": 50, "release_name": "Employment Situation", "date": "2026-10-02"},
            {"release_id": 999, "release_name": "Other", "date": "2026-10-14"},
        ]}
        response = mock.Mock(json=mock.Mock(return_value=payload), raise_for_status=mock.Mock())
        with mock.patch.object(events.requests, "get", return_value=response) as get:
            result = events.fred_events("k", date(2026, 10, 1), date(2026, 10, 15))
        self.assertEqual(get.call_args.kwargs["params"]["include_release_dates_with_no_data"], "true")
        self.assertEqual([(e.day, e.time_text, e.title) for e in result], [
            (date(2026, 10, 14), "20:30", "CPI 消費者物價指數"),
            (date(2026, 10, 2), "20:30", "非農就業報告"),
        ])

    def test_collect_events_window_and_failures(self):
        now = datetime(2026, 10, 14, 21, 0, tzinfo=TW)
        cpi_past = events._timed_event(date(2026, 10, 14), events.FRED_RELEASE_TIME, "CPI", 3)
        ppi = events._timed_event(date(2026, 10, 15), events.FRED_RELEASE_TIME, "PPI", 2)
        far = events._timed_event(date(2026, 10, 30), events.FRED_RELEASE_TIME, "GDP", 2)
        warnings = []
        with mock.patch.dict(os.environ, {"FRED_API_KEY": "k", "ALPHAVANTAGE_API_KEY": "a"}), \
                mock.patch.object(events, "fred_events", return_value=[cpi_past, ppi, far]), \
                mock.patch.object(events, "earnings_events", side_effect=ValueError("額度用完")):
            result = events.collect_events(now, warn=warnings.append)
        titles = [e.title for e in result]
        self.assertNotIn("CPI", titles)  # 已經公布
        self.assertNotIn("GDP", titles)  # 超過 7 天
        self.assertIn("PPI", titles)
        self.assertIn("美股選擇權月到期", titles)  # 10/16
        self.assertIn("台指期／台指選擇權月結算", titles)  # 10/21
        self.assertEqual([e.day for e in result], sorted(e.day for e in result))
        self.assertTrue(any("財報資料讀取失敗" in w for w in warnings))

    def test_missing_keys_warn(self):
        warnings = []
        events.collect_events(datetime(2026, 10, 14, 18, 0, tzinfo=TW), warn=warnings.append)
        self.assertTrue(any("FRED_API_KEY" in w for w in warnings))
        self.assertTrue(any("ALPHAVANTAGE_API_KEY" in w for w in warnings))

    def test_section_and_summary_lines(self):
        ev = [events._timed_event(date(2026, 10, 15), events.FRED_RELEASE_TIME, "PPI 生產者物價指數", 2),
              events.Event(date(2026, 10, 16), "", "美股選擇權月到期", 2)]
        lines = events.event_lines(events.build_events_section(ev))
        self.assertEqual(lines[1:], [
            "【未來 7 天重大事件】",
            "10/15（四） 20:30 PPI 生產者物價指數 ⭐⭐",
            "10/16（五） 美股選擇權月到期 ⭐⭐",
        ])


class PagesTest(unittest.TestCase):
    def test_publish_builds_index(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports, site = Path(tmp) / "reports", Path(tmp) / "site"
            reports.mkdir()
            for day in ("2026-09-24", "2026-09-25"):
                (reports / f"chip_report_{day}.html").write_text(
                    f"<html><body>\n<main>\n<h1>{day}</h1></main></body></html>", encoding="utf-8"
                )
                self.assertEqual(pages.publish(site, reports / f"chip_report_{day}.html"), day)
            self.assertTrue((site / ".nojekyll").exists())
            index = (site / "index.html").read_text(encoding="utf-8")
            self.assertLess(index.index("2026-09-25.html"), index.index("2026-09-24.html"))
            self.assertIn("2026-09-25（五）", index)
            self.assertIn("共 2 份日報、0 份週報", index)
            self.assertNotIn("每週週報", index)
            page = (site / "2026-09-25.html").read_text(encoding="utf-8")
            self.assertIn('href="index.html"', page)

    def test_publish_weekly(self):
        with tempfile.TemporaryDirectory() as tmp:
            site = Path(tmp) / "site"
            for name in ("chip_report_2026-09-25.html", "weekly_report_2026-09-25.html"):
                src = Path(tmp) / name
                src.write_text("<main>\n</main>", encoding="utf-8")
                self.assertEqual(pages.publish(site, src), "2026-09-25")
            self.assertTrue((site / "weekly-2026-09-25.html").exists())
            index = (site / "index.html").read_text(encoding="utf-8")
            self.assertIn('href="weekly-2026-09-25.html"><span>09/21～09/25 週報</span>', index)
            self.assertLess(index.index("每週週報"), index.index("每日報告"))
            self.assertIn("共 1 份日報、1 份週報", index)

    def test_publish_rejects_bad_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "report.html"
            bad.write_text("x", encoding="utf-8")
            with self.assertRaises(ValueError):
                pages.publish(Path(tmp) / "site", bad)


class NotifyCliTest(unittest.TestCase):
    def test_sends_summary_file(self):
        env = {"LINE_CHANNEL_ACCESS_TOKEN": "ln", "LINE_USER_ID": "U1",
               "TELEGRAM_BOT_TOKEN": "", "TELEGRAM_CHAT_ID": ""}
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.dict(os.environ, env), \
                mock.patch("notify.requests.post", return_value=mock.Mock(status_code=200)) as post, \
                mock.patch("sys.stdout"):
            path = Path(tmp) / "summary.txt"
            path.write_text("摘要內容", encoding="utf-8")
            self.assertEqual(notify.main(["notify.py", str(path)]), 0)
        self.assertEqual(post.call_args.kwargs["json"]["messages"][0]["text"], "摘要內容")

    def _send_with_env(self, env):
        base = {"LINE_CHANNEL_ACCESS_TOKEN": "ln", "LINE_USER_ID": "Uself", "LINE_USER_IDS": "",
                "LINE_SEND_MODE": "", "TELEGRAM_BOT_TOKEN": "", "TELEGRAM_CHAT_ID": ""}
        base.update(env)
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.dict(os.environ, base), \
                mock.patch("notify.requests.post", return_value=mock.Mock(status_code=200)) as post, \
                mock.patch("sys.stdout") as out, mock.patch("sys.stderr"):
            path = Path(tmp) / "summary.txt"
            path.write_text("日報\f週報", encoding="utf-8")
            code = notify.main(["notify.py", str(path)])
        printed = "".join(c.args[0] for c in out.write.call_args_list)
        return code, post, printed

    def test_line_mode_self_default(self):
        code, post, printed = self._send_with_env({})
        self.assertEqual(code, 0)
        self.assertEqual(post.call_args.args[0], "https://api.line.me/v2/bot/message/push")
        self.assertEqual(post.call_args.kwargs["json"]["to"], "Uself")
        self.assertEqual(len(post.call_args.kwargs["json"]["messages"]), 2)

    def test_line_mode_broadcast(self):
        code, post, printed = self._send_with_env({"LINE_SEND_MODE": "Broadcast"})
        self.assertEqual(code, 0)
        self.assertEqual(post.call_args.args[0], "https://api.line.me/v2/bot/message/broadcast")
        self.assertNotIn("to", post.call_args.kwargs["json"])
        self.assertIn("廣播給所有好友", printed)

    def test_line_mode_list(self):
        ids = "Uself #我\nUfriend1 #小明, Ufriend2\n\nUfriend1"
        code, post, printed = self._send_with_env({"LINE_SEND_MODE": "list", "LINE_USER_IDS": ids})
        self.assertEqual(code, 0)
        self.assertEqual(post.call_args.args[0], "https://api.line.me/v2/bot/message/multicast")
        self.assertEqual(post.call_args.kwargs["json"]["to"], ["Uself", "Ufriend1", "Ufriend2"])
        self.assertIn("3 位指定對象", printed)

    def test_line_mode_invalid_or_empty_list(self):
        self.assertEqual(self._send_with_env({"LINE_SEND_MODE": "all"})[0], 1)
        code, post, _ = self._send_with_env({"LINE_SEND_MODE": "list"})
        self.assertEqual(code, 1)
        post.assert_not_called()

    def test_failure_returns_1(self):
        env = {"LINE_CHANNEL_ACCESS_TOKEN": "ln", "LINE_USER_ID": "U1"}
        bad = mock.Mock(status_code=400, text="bad")
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.dict(os.environ, env), \
                mock.patch("notify.requests.post", return_value=bad), \
                mock.patch("sys.stderr"):
            path = Path(tmp) / "summary.txt"
            path.write_text("x", encoding="utf-8")
            self.assertEqual(notify.main(["notify.py", str(path)]), 1)


class MainTest(unittest.TestCase):
    def setUp(self):
        # 固定「今天」，讓自動找交易日的結果不受實際日期影響
        patcher = mock.patch.object(chip_report, "taipei_today", return_value=date(2026, 9, 29))
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_main_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, {"FINMIND_TOKEN": "t"}), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--output-dir", tmp]), \
                mock.patch("sys.stdout"):
            self.assertEqual(chip_report.main(), 0)
            names = sorted(p.name for p in Path(tmp).iterdir())
        # 2026-09-25 是週五，會一併產生週報
        self.assertEqual(
            names,
            ["chip_report_2026-09-25.html", "chip_report_2026-09-25.md", "chip_report_2026-09-25.xlsx",
             "weekly_report_2026-09-25.html", "weekly_report_2026-09-25.md",
             "weekly_report_2026-09-25.xlsx"],
        )

    def test_summary_file_with_report_url(self):
        env = {"FINMIND_TOKEN": "t", "REPORT_BASE_URL": "https://u.github.io/repo/"}
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, env), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--output-dir", tmp,
                                                "--summary-file", f"{tmp}/summary.txt"]), \
                mock.patch("sys.stdout"):
            self.assertEqual(chip_report.main(), 0)
            summary = Path(tmp, "summary.txt").read_text(encoding="utf-8")
        self.assertIn("【外資連買】", summary)
        self.assertIn("完整報告：https://u.github.io/repo/2026-09-25.html", summary)
        self.assertIn("外資台指期 淨未平倉 -30,000 口（日增減 -2,000）", summary)
        self.assertIn("P/C Ratio 120.00%（前日 100.00%）", summary)
        self.assertIn("00878 國泰永續高股息 +1,500 張", summary)
        self.assertIn("1. 2330 台積電 +1,000 張", summary)
        self.assertIn("1. 2330 台積電 +0.50 百分點", summary)

    def test_etfs_excluded_from_stock_rankings(self):
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, {"FINMIND_TOKEN": "t"}), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--output-dir", tmp,
                                                "--formats", "md"]), \
                mock.patch("sys.stdout"):
            self.assertEqual(chip_report.main(), 0)
            md = Path(tmp, "chip_report_2026-09-25.md").read_text(encoding="utf-8")
        before_etf = md.split("## 高股息 ETF")[0]
        self.assertNotIn("00878", before_etf)  # 不在個股排行
        self.assertIn("| 00878 | 國泰永續高股息 |", md)
        headings = [line for line in md.splitlines() if line.startswith("## ")]
        self.assertEqual(headings[:2], ["## 全市場總覽", "## 期貨選擇權籌碼"])

    def test_extra_failure_does_not_break_report(self):
        def failing_fetch(token, dataset, **params):
            if dataset in ("TaiwanStockMarginPurchaseShortSale", "TaiwanFuturesInstitutionalInvestors"):
                raise finmind.FinMindError("權限不足")
            return fake_fetch(token, dataset, **params)

        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(finmind, "fetch_dataset", failing_fetch), \
                mock.patch.dict(os.environ, {"FINMIND_TOKEN": "t"}), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--output-dir", tmp,
                                                "--formats", "md"]), \
                mock.patch("sys.stdout"), mock.patch("sys.stderr") as err:
            self.assertEqual(chip_report.main(), 0)
            md = Path(tmp, "chip_report_2026-09-25.md").read_text(encoding="utf-8")
        self.assertNotIn("融資增加排行", md)
        self.assertNotIn("期貨選擇權籌碼", md)
        self.assertIn("千張大戶持股增加", md)
        self.assertIn("外資買超排行", md)
        warnings = "".join(c.args[0] for c in err.write.call_args_list)
        self.assertIn("融資融券資料讀取失敗", warnings)

    def test_missing_token(self):
        with mock.patch.dict(os.environ, {"FINMIND_TOKEN": ""}), \
                mock.patch.object(sys, "argv", ["chip_report.py"]), \
                mock.patch("sys.stderr"):
            self.assertEqual(chip_report.main(), 1)

    def test_notify_without_channels(self):
        env = {"FINMIND_TOKEN": "t", "TELEGRAM_BOT_TOKEN": "", "LINE_CHANNEL_ACCESS_TOKEN": ""}
        with mock.patch.dict(os.environ, env), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--notify"]), \
                mock.patch("sys.stderr"):
            self.assertEqual(chip_report.main(), 1)

    def test_notify_sends_to_both_channels(self):
        env = {
            "FINMIND_TOKEN": "t",
            "TELEGRAM_BOT_TOKEN": "tg", "TELEGRAM_CHAT_ID": "123",
            "LINE_CHANNEL_ACCESS_TOKEN": "ln", "LINE_USER_ID": "U1",
        }
        ok = mock.Mock(status_code=200)
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, env), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--notify", "--formats", "md",
                                                "--output-dir", tmp]), \
                mock.patch("notify.requests.post", return_value=ok) as post, \
                mock.patch("sys.stdout"):
            self.assertEqual(chip_report.main(), 0)
        urls = [c.args[0] for c in post.call_args_list]
        # 週五：Telegram 分兩則（日報、週報），LINE 同一次推播兩個對話框
        self.assertEqual(urls, ["https://api.telegram.org/bottg/sendMessage"] * 2
                         + ["https://api.line.me/v2/bot/message/push"])
        line_body = post.call_args_list[2].kwargs["json"]
        self.assertEqual(line_body["to"], "U1")
        self.assertEqual(len(line_body["messages"]), 2)
        self.assertIn("【外資買超】", line_body["messages"][0]["text"])
        self.assertIn("台股每週籌碼週報", line_body["messages"][1]["text"])

    def _run_today_only(self, today):
        env = {"FINMIND_TOKEN": "t", "TELEGRAM_BOT_TOKEN": "tg", "TELEGRAM_CHAT_ID": "1"}
        ok = mock.Mock(status_code=200)
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(chip_report, "taipei_today", return_value=today), \
                mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, env), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--today-only", "--notify",
                                                "--output-dir", tmp]), \
                mock.patch("notify.requests.post", return_value=ok) as post, \
                mock.patch("sys.stdout"):
            code = chip_report.main()
            files = list(Path(tmp).iterdir())
        return code, files, post

    def test_today_only_with_data(self):
        code, files, post = self._run_today_only(date(2026, 9, 25))
        self.assertEqual(code, 0)
        self.assertEqual(len(files), 6)  # 日報＋週報
        self.assertEqual(post.call_count, 2)  # Telegram 日報、週報各一則

    def test_today_only_without_data_skips(self):
        code, files, post = self._run_today_only(date(2026, 9, 28))
        self.assertEqual(code, 0)
        self.assertEqual(files, [])
        post.assert_not_called()

    def test_today_only_with_date_skips_when_no_data(self):
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, {"FINMIND_TOKEN": "t"}), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--date", "2026-09-28", "--today-only",
                                                "--output-dir", tmp]), \
                mock.patch("sys.stdout"):
            self.assertEqual(chip_report.main(), 0)
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_date_without_today_only_still_errors_when_no_data(self):
        with mock.patch.object(finmind, "fetch_dataset", fake_fetch), \
                mock.patch.dict(os.environ, {"FINMIND_TOKEN": "t"}), \
                mock.patch.object(sys, "argv", ["chip_report.py", "--date", "2026-09-28"]), \
                mock.patch("sys.stdout"), mock.patch("sys.stderr"):
            self.assertEqual(chip_report.main(), 1)

    def test_taipei_today_uses_utc_plus_8(self):
        fake_now = chip_report.datetime(2026, 9, 28, 17, 0, tzinfo=chip_report.timezone.utc)
        with mock.patch.object(chip_report, "datetime") as dt:
            dt.now.side_effect = lambda tz: fake_now.astimezone(tz)
            self.assertEqual(chip_report.taipei_today(), date(2026, 9, 29))


if __name__ == "__main__":
    unittest.main()
