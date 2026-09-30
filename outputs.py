"""報告輸出：Markdown、HTML、Excel"""

import html

from analysis import report_title

SUBTITLE = "資料來源：FinMind（上市＋上櫃普通股；外資含外資自營商）"


def format_cell(value, kind):
    if value is None:
        return "-"
    if kind == "text":
        return str(value)
    if kind == "int":
        return f"{value:,}"
    if kind == "signed_int":
        return f"{value:+,}"
    if kind == "price":
        return f"{value:,.2f}"
    if kind == "pct":
        return f"{value:.2f}%"
    if kind == "signed_pct":
        return f"{value:+.2f}%"
    if kind in ("signed_yi", "signed_float"):
        return f"{value:+,.2f}"
    raise ValueError(f"未知的欄位格式：{kind}")


def _report_sections(report):
    return [report.overview] + report.sections


# ---------- Markdown ----------

def _md_align(kind):
    return ":---" if kind == "text" else "---:"


def section_to_markdown(section):
    lines = [f"## {section.title}", ""]
    if section.note:
        lines += [f"> {section.note}", ""]
    if not section.rows:
        lines += ["（無資料）", ""]
        return "\n".join(lines)
    lines.append("| " + " | ".join(c.header for c in section.columns) + " |")
    lines.append("|" + "|".join(_md_align(c.kind) for c in section.columns) + "|")
    for row in section.rows:
        cells = [format_cell(v, c.kind) for v, c in zip(row, section.columns)]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    return "\n".join(lines)


def render_markdown(report):
    parts = [f"# {report_title(report)}", "", SUBTITLE, ""]
    parts += [section_to_markdown(s) for s in _report_sections(report)]
    return "\n".join(parts)


# ---------- HTML ----------

HTML_STYLE = """
:root {
  --bg: #f6f7f9; --card: #ffffff; --text: #1f2328; --muted: #656d76;
  --border: #d8dee4; --head: #eef1f4; --up: #d1242f; --down: #1a7f37;
}
@media (prefers-color-scheme: dark) {
  :root {
    --bg: #0d1117; --card: #161b22; --text: #e6edf3; --muted: #8d96a0;
    --border: #30363d; --head: #1f252d; --up: #ff6b6b; --down: #3fb950;
  }
}
* { box-sizing: border-box; }
body {
  margin: 0; padding: 16px; background: var(--bg); color: var(--text);
  font-family: -apple-system, "Segoe UI", "Microsoft JhengHei", "PingFang TC", sans-serif;
  font-size: 14px; line-height: 1.5;
}
main { max-width: 1100px; margin: 0 auto; }
h1 { font-size: 22px; margin: 0 0 4px; }
.subtitle { color: var(--muted); margin: 0 0 16px; }
nav { display: flex; flex-wrap: wrap; gap: 8px; margin-bottom: 16px; }
nav a {
  color: var(--text); text-decoration: none; background: var(--card);
  border: 1px solid var(--border); border-radius: 999px; padding: 4px 12px; font-size: 13px;
}
section {
  background: var(--card); border: 1px solid var(--border); border-radius: 8px;
  padding: 12px 16px; margin-bottom: 16px;
}
h2 { font-size: 17px; margin: 0 0 8px; }
.note { color: var(--muted); font-size: 12px; margin: 0 0 8px; }
.table-wrap { overflow-x: auto; }
table { border-collapse: collapse; width: 100%; white-space: nowrap; }
th, td { padding: 6px 10px; border-bottom: 1px solid var(--border); }
th { background: var(--head); font-weight: 600; text-align: left; }
td.num, th.num { text-align: right; font-variant-numeric: tabular-nums; }
td.up { color: var(--up); }
td.down { color: var(--down); }
tbody tr:hover { background: var(--head); }
"""


def _html_cell(value, kind):
    classes = []
    if kind != "text":
        classes.append("num")
    if kind.startswith("signed_") and value is not None:
        if value > 0:
            classes.append("up")
        elif value < 0:
            classes.append("down")
    cls = f' class="{" ".join(classes)}"' if classes else ""
    return f"<td{cls}>{html.escape(format_cell(value, kind))}</td>"


def section_to_html(section):
    out = [f'<section id="{section.key}">', f"<h2>{html.escape(section.title)}</h2>"]
    if section.note:
        out.append(f'<p class="note">{html.escape(section.note)}</p>')
    if not section.rows:
        out.append("<p>（無資料）</p></section>")
        return "\n".join(out)
    out.append('<div class="table-wrap"><table><thead><tr>')
    for c in section.columns:
        cls = "" if c.kind == "text" else ' class="num"'
        out.append(f"<th{cls}>{html.escape(c.header)}</th>")
    out.append("</tr></thead><tbody>")
    for row in section.rows:
        cells = "".join(_html_cell(v, c.kind) for v, c in zip(row, section.columns))
        out.append(f"<tr>{cells}</tr>")
    out.append("</tbody></table></div></section>")
    return "\n".join(out)


def render_html(report):
    title = html.escape(report_title(report))
    sections = _report_sections(report)
    nav = "".join(
        f'<a href="#{s.key}">{html.escape(s.title.split("（")[0])}</a>' for s in sections
    )
    body = "\n".join(section_to_html(s) for s in sections)
    return f"""<!doctype html>
<html lang="zh-Hant-TW">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>{HTML_STYLE}</style>
</head>
<body>
<main>
<h1>{title}</h1>
<p class="subtitle">{html.escape(SUBTITLE)}｜紅色為買超／上漲，綠色為賣超／下跌</p>
<nav>{nav}</nav>
{body}
</main>
</body>
</html>
"""


# ---------- Excel ----------

EXCEL_FORMATS = {
    "text": "@",
    "int": "#,##0",
    "signed_int": "+#,##0;-#,##0;0",
    "price": "#,##0.00",
    "pct": '0.00"%"',
    "signed_pct": '+0.00"%";-0.00"%";0.00"%"',
    "signed_yi": "+#,##0.00;-#,##0.00;0.00",
    "signed_float": "+0.00;-0.00;0.00",
}

SHEET_NAMES = {
    "overview": "總覽",
    "foreign_buy": "外資買超",
    "foreign_sell": "外資賣超",
    "trust_buy": "投信買超",
    "trust_sell": "投信賣超",
    "sync_buy": "外資投信同步買超",
    "foreign_streak": "外資連續買超",
    "trust_streak": "投信連續買超",
    "industry": "產業別",
    "etf": "高股息ETF",
    "futures": "期貨選擇權",
    "margin_up": "融資增加",
    "margin_down": "融資減少",
    "short_up": "融券增加",
    "holder_up": "千張大戶增加",
    "holder_down": "千張大戶減少",
    "all": "全部個股",
}


def write_excel(report, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="EEF1F4")
    up_font = Font(color="D1242F")
    down_font = Font(color="1A7F37")

    wb = Workbook()
    wb.remove(wb.active)
    for section in _report_sections(report) + [report.all_stocks]:
        ws = wb.create_sheet(SHEET_NAMES.get(section.key, section.key)[:31])
        for col_idx, column in enumerate(section.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=column.header)
            cell.font = header_font
            cell.fill = header_fill
            width = max(len(column.header) * 2 + 2, 10)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        for row_idx, row in enumerate(section.rows, 2):
            for col_idx, (value, column) in enumerate(zip(row, section.columns), 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.number_format = EXCEL_FORMATS[column.kind]
                if column.kind.startswith("signed_") and value is not None:
                    if value > 0:
                        cell.font = up_font
                    elif value < 0:
                        cell.font = down_font
        ws.freeze_panes = "A2"
        if section.rows:
            ws.auto_filter.ref = ws.dimensions
        if section.note:
            ws.cell(row=len(section.rows) + 3, column=1, value=section.note)
    wb.save(path)
